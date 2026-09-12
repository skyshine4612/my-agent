# app/skills/excel/scripts/inspect.py
# 看 Excel 表结构：sheet 清单 + 每个 sheet 的列名、样本值、行数。
# 样本值是刻意返回的：用户说的「手机号」在表里可能叫「联系电话」，只给列名模型只能猜，
# 给两个样本值它就能认出来。截断上限是防宽表（几十列、长文本单元格）撑爆模型上下文。
# 沙箱内执行：只 import 标准库 + openpyxl，不能 import 项目代码。
import itertools
import json
import os
import sys

MAX_COLUMNS = 30
PREVIEW_ROWS = 5
SAMPLE_ROWS = 3
SAMPLE_CHARS = 32


def _text(value):
    """单元格值转成短的、可读的字符串：None 当空串，长文本截断。"""
    if value is None:
        return ""
    s = str(value).strip()
    return f"{s[:SAMPLE_CHARS]}…" if len(s) > SAMPLE_CHARS else s


def _raw(value):
    """参数用字符串：只去空白。

    **不能走 _text**：那个带 SAMPLE_CHARS 显示截断，而 file_id 加扩展名有 37 字符，
    截一刀就变成另一个不存在的文件名。
    """
    return "" if value is None else str(value).strip()


def _resolve(args, work_dir):
    """定位输入文件：后端给的是沙箱内本次执行的 work_dir 下的文件名，这里兼容绝对路径。"""
    ref = _raw(args.get("file_id"))
    if not ref:
        raise SystemExit("缺少 file_id：把会话文件清单里的 file_id 原样填进来")
    return ref if os.path.isabs(ref) else os.path.join(work_dir, os.path.basename(ref))


def summarize_sheet(name, preview_rows, total_rows, header_row=1):
    """整理一个 sheet 的结构摘要（纯函数，便于单测）。

    preview_rows 是 sheet 开头若干行（含表头行本身，行号从 1 算起），total_rows 是整表行数。
    """
    head = max(0, header_row - 1)
    header = [_text(c) for c in preview_rows[head]] if head < len(preview_rows) else []
    # 空表头列用位置兜底命名：脱敏要按列名指定目标列，无名列也得有个能指认的名字
    columns = [h or f"第{i + 1}列" for i, h in enumerate(header)][:MAX_COLUMNS]
    samples = {}
    for i, col in enumerate(columns):
        values = []
        for row in preview_rows[head + 1:]:
            v = _text(row[i]) if i < len(row) else ""
            if v and v not in values:
                values.append(v)
            if len(values) >= SAMPLE_ROWS:
                break
        samples[col] = values
    return {"name": name, "row_count": total_rows, "header_row": header_row, "columns": columns,
            "samples": samples, "preview": [[_text(c) for c in row] for row in preview_rows]}


def main():
    payload = json.load(sys.stdin)
    args, work_dir = payload["args"], payload["work_dir"]

    from openpyxl import load_workbook  # 延迟导入：纯逻辑单测不必装 openpyxl

    path = _resolve(args, work_dir)
    if not os.path.exists(path):
        raise SystemExit(f"读不到文件：{path}")

    header_row = int(args.get("header_row") or 1)
    # read_only：只读结构，不把整表载入内存（宽表能到几十万行）
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = [summarize_sheet(ws.title,
                              [[c.value for c in row] for row in itertools.islice(ws.iter_rows(), PREVIEW_ROWS)],
                              ws.max_row or 0, header_row)
              for ws in wb.worksheets]
    wb.close()
    print(json.dumps({"sheets": sheets}, ensure_ascii=False))


if __name__ == "__main__":
    main()
