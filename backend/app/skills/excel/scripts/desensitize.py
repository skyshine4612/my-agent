# app/skills/excel/scripts/desensitize.py
# Excel 字段脱敏：按 ops 逐列套规则，产出一份新文件（绝不覆盖原件）。
# 沙箱内执行：只 import 标准库 + openpyxl，不能 import 项目代码（沙箱里没有项目代码）。
# 纯逻辑（apply_rules）与 IO（main）分开：脱敏语义能被单测覆盖，不用起容器。
import functools
import hashlib
import json
import os
import sys


def _text(value):
    """单元格值转成去空白的字符串：None 当空串。"""
    return "" if value is None else str(value).strip()


def _rule(fn):
    """把「字符串 → 字符串」的规则包成「单元格值 → 单元格值」：空值原样返回，不把 None 变成空串。"""

    @functools.wraps(fn)
    def wrapped(value, salt):
        s = _text(value)
        return value if not s else fn(s, salt)

    return wrapped


@_rule
def _phone(s, salt):
    """11 位手机号打码中间 4 位；位数不对（座机号、「无」）原样保留。"""
    return f"{s[:3]}****{s[-4:]}" if len(s) == 11 else s


@_rule
def _name(s, salt):
    """姓名只留姓。"""
    return s[0] + "*" * (len(s) - 1)


@_rule
def _idcard(s, salt):
    """身份证留前 6 后 4；填了半截、位数不够的原样保留。"""
    return f"{s[:6]}{'*' * (len(s) - 10)}{s[-4:]}" if len(s) >= 10 else s


@_rule
def _email(s, salt):
    """邮箱只留首字母：zhangsan@qq.com → z***@qq.com；不含 @ 的原样保留。"""
    local, at, domain = s.partition("@")
    return f"{local[0]}***@{domain}" if at and local else s


@_rule
def _bankcard(s, salt):
    """银行卡只留末 4 位。"""
    return f"{'*' * (len(s) - 4)}{s[-4:]}" if len(s) > 4 else s


@_rule
def _hash(s, salt):
    """加盐哈希前 16 位：同值同结果（保住可关联性），不可反查。"""
    return hashlib.sha256(f"{salt}{s}".encode("utf-8")).hexdigest()[:16]


# 规则表：键是给模型看的规则名；drop 没有函数，在 _plan 里单独处理
RULES = {
    "phone_middle4": _phone,
    "name_keep_first": _name,
    "idcard_keep_ends": _idcard,
    "email_mask": _email,
    "bankcard_keep_last4": _bankcard,
    "hash": _hash,
    "drop": None,
}


def _plan(header, ops):
    """把 ops 解析成 (要删的列索引, {列索引: 规则函数}, 规范化后的表头)。

    列名或规则名对不上就抛 ValueError —— 脚本自己校验参数，给一句模型能照着改的话，
    比甩栈回溯有用得多。
    """
    if not ops:
        raise ValueError("ops 不能为空：请说明要脱敏哪些列")
    header = [_text(h) for h in header]
    drop, rule_of = [], {}
    for op in ops:
        col, rule = _text(op.get("column")), op.get("rule")
        if rule not in RULES:
            raise ValueError(f"不认识的规则：{rule}（可选：{'、'.join(RULES)}）")
        if col not in header:
            raise ValueError(f"表里没有这一列：{col}（实际列名：{'、'.join(h for h in header if h)}）")
        index = header.index(col)
        if rule == "drop":
            drop.append(index)
        else:
            rule_of[index] = RULES[rule]
    return drop, rule_of, header


def apply_rules(header, rows, ops, salt=""):
    """按 ops 算出脱敏结果，返回 (新表头, 新数据行, 保留的原始列索引)。"""
    drop, rule_of, header = _plan(header, ops)
    keep = [i for i in range(len(header)) if i not in drop]
    new_rows = [[rule_of[i](row[i], salt) if i in rule_of else (row[i] if i < len(row) else None)
                 for i in keep] for row in rows]
    return [header[i] for i in keep], new_rows, keep


def _resolve(args, work_dir):
    """定位输入文件：后端已把 file_id 换成沙箱内绝对路径，这里兼容直接给裸文件名的情况。"""
    ref = _text(args.get("file_id"))
    if not ref:
        raise SystemExit("缺少 file_id：把会话文件清单里的 file_id 原样填进来")
    return ref if os.path.isabs(ref) else os.path.join(work_dir, os.path.basename(ref))


def _safe_name(name):
    """产物一律落在 out_dir 里：剥掉目录部分，且必须是 .xlsx。"""
    name = os.path.basename(_text(name)) or "已脱敏.xlsx"
    return name if name.lower().endswith(".xlsx") else f"{name}.xlsx"


def main():
    payload = json.load(sys.stdin)
    args, work_dir, out_dir = payload["args"], payload["work_dir"], payload["out_dir"]

    from openpyxl import load_workbook  # 延迟导入：纯逻辑单测不必装 openpyxl

    path = _resolve(args, work_dir)
    if not os.path.exists(path):
        raise SystemExit(f"读不到文件：{path}")

    # 不用 read_only：写回时要保留原有样式、列宽与合并单元格
    wb = load_workbook(path)
    sheet_name = _text(args.get("sheet")) or wb.sheetnames[0]
    if sheet_name not in wb.sheetnames:
        raise SystemExit(f"没有这个 sheet：{sheet_name}（实际有：{'、'.join(wb.sheetnames)}）")
    ws = wb[sheet_name]

    header_row = int(args.get("header_row") or 1)
    header = [c.value for c in ws[header_row]]
    data_rows = [[c.value for c in r] for r in ws.iter_rows(min_row=header_row + 1)]

    try:
        new_header, new_rows, keep = apply_rules(header, data_rows, args.get("ops") or [], args.get("salt", ""))
    except ValueError as e:
        raise SystemExit(str(e))

    # 先删列（从右往左，避免删前面的列把后面挪位），再把新值逐格写回保留的列
    for i in sorted(set(range(len(header))) - set(keep), reverse=True):
        ws.delete_cols(i + 1)
    for new_row_no, row in enumerate(new_rows, start=header_row + 1):
        for pos, value in enumerate(row):
            ws.cell(row=new_row_no, column=pos + 1).value = value

    out_name = _safe_name(args.get("out_name") or "已脱敏.xlsx")
    wb.save(os.path.join(out_dir, out_name))
    print(json.dumps({
        "sheet": sheet_name,
        "rows": len(new_rows),
        "columns": new_header,
        "dropped": [header[i] for i in sorted(set(range(len(header))) - set(keep))],
        "out_name": out_name,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
